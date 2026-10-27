# -*- coding: utf-8 -*-
"""导入便捷性/准确性验证夹具生成器（确定性输出）。

Run: python tests/fixtures/import-ux/build_fixtures.py
产出（与脚本同目录）：
  1. 现场检测记录-乱序列.csv   事件台账：别名表头 + 乱序列 + 脏数据 + BOM
  2. 受热面材料台账-别名表头.csv 材料库：别名表头 + 重复行 + 缺材质
  3. 无关表格-考勤.csv          非台账表格：必须判为 unknown、不误判
  4. 检测记录.xlsx              真实 Excel（openpyxl）：多工作表，首个为乱序测厚记录
  5. 检测台账.docx              真实 Word（zipfile + OOXML w:tbl）：别名表头测厚记录
"""
import csv
import os
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))


def write_csv(name, rows, bom=True):
    path = os.path.join(HERE, name)
    with open(path, "w", encoding="utf-8-sig" if bom else "utf-8", newline="") as fh:
        csv.writer(fh).writerows(rows)
    print("written", path)


def build_events_csv():
    # 表头为现场习惯叫法，顺序与系统模板不同；列含义靠智能匹配对齐
    header = ["设备编码", "机组号", "检测日期", "分类", "钢管规格", "材质牌号",
              "实测壁厚值(mm)", "布氏硬度值", "检测类型", "结论备注", "新管段编码"]
    rows = [
        header,
        ["1-WW-FR-008-0003-000", "1", "2024年3月5日", "事件", "Φ63.5×7.5", "SA-210C", "6.8", "182", "测厚", "前墙第8排第3管", ""],
        ["1-WW-FR-008-0004-000", "1", "2024/4/8", "事件", "Φ63.5×7.5", "SA-210C", "7.1", "180", "测厚", "目视无异常", ""],
        ["1-WW-FR-020-0010-000", "1", "2024.4.15", "事件", "Φ63.5×7.5", "SA-210C", "7.05", "180", "金相分析", "组织正常", ""],
        ["1-WW-FR-080-0003-000", "1", "2024-05-06", "事件", "Φ63.5×7.5", "SA-210C", "6.9", "181", "测厚", "屏号待核对", ""],
        ["1-WW-FR-008-0003-000", "1", "2024年3月5日", "事件", "Φ63.5×7.5", "SA-210C", "6.8", "182", "测厚", "前墙第8排第3管", ""],
        ["", "1", "2024-06-01", "事件", "Φ63.5×7.5", "SA-210C", "7.0", "175", "测厚", "缺编码", ""],
        ["1-WW-FR-022-0012-000", "2", "2024-02-11", "事件", "Φ63.5×7.5", "SA-210C", "7.01", "176", "测厚", "炉号冲突", ""],
        ["1-WW-FR-021-0011-000", "1", "2024年13月45日", "事件", "Φ63.5×7.5", "SA-210C", "7.02", "178", "测厚", "日期待核对", ""],
        ["1-WW-FR-021-0012-000", "1", "2025-01-20", "事件", "Φ63.5×7.5", "SA-210C", "6.5~7.0", "179", "测厚", "厚度为区间值", ""],
        ["1-WW-FL-010-0005-000", "1", "2026-05-22", "更换", "Φ60×9", "SA-210C", "", "", "爆管更换", "", "1-WW-FL-010-0005-000-R1"],
        [""] * len(header),  # 空行：应被自动丢弃
    ]
    write_csv("现场检测记录-乱序列.csv", rows)


def build_material_csv():
    header = ["部件", "系统", "位置", "结构", "规格", "管径(mm)", "壁厚(mm)",
              "材质牌号", "库存数量", "总长度", "预留", "说明"]
    rows = [
        header,
        ["省煤器", "ECO", "入口第1排", "蛇形管", "Φ51×6", "51", "6", "SA-210C", "10", "60", "A区", "第47页"],
        ["水冷壁", "WW", "前墙第8排", "膜式壁", "Φ63.5×7.5", "63.5", "7.5", "SA-210C", "4", "24", "B区", "备用"],
        ["省煤器", "ECO", "入口第1排", "蛇形管", "Φ51×6", "51", "6", "SA-210C", "10", "60", "A区", "第47页"],
        ["高温过热器", "HSH", "出口第2排", "吊管", "Φ51×8", "51", "8", "", "", "", "", "缺材质"],
    ]
    write_csv("受热面材料台账-别名表头.csv", rows)


def build_unrelated_csv():
    rows = [
        ["姓名", "部门", "日期", "工时", "备注"],
        ["张三", "检修部", "2024-03-05", "8", "夜间抢修"],
        ["李四", "检测中心", "2024-03-06", "7.5", ""],
    ]
    write_csv("无关表格-考勤.csv", rows, bom=False)


def build_xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "测厚记录"
    ws.append(["设备编码", "机组号", "检测日期", "分类", "实测壁厚值(mm)", "布氏硬度值", "检测类型", "结论备注"])
    ws.append(["1-WW-FR-008-0003-000", "1", "2024年3月5日", "事件", "6.8", "182", "测厚", "前墙第8排第3管"])
    ws.append(["1-WW-FR-008-0004-000", "1", "2024/4/8", "事件", "7.1", "180", "测厚", "目视无异常"])
    ws.append(["1-WW-FR-008-0004-000", "1", "2024/4/8", "事件", "7.1", "180", "测厚", "目视无异常"])
    ws.append(["1-WW-FR-021-0011-000", "1", "2024年13月45日", "事件", "7.02", "178", "测厚", "日期待核对"])
    ws2 = wb.create_sheet("材料库")
    ws2.append(["部件", "系统", "规格", "材质", "库存数量"])
    ws2.append(["省煤器", "ECO", "Φ51×6", "SA-210C", "10"])
    path = os.path.join(HERE, "检测记录.xlsx")
    wb.save(path)
    print("written", path)


def cell(text):
    return ("<w:tc><w:p><w:pPr><w:jc w:val=\"center\"/></w:pPr>"
            "<w:r><w:t xml:space=\"preserve\">" + text + "</w:t></w:r></w:p></w:tc>")


def build_docx():
    header = ["管段编号", "检测时间", "类别", "壁厚值", "硬度值", "备注"]
    data = [
        ["1-WW-FR-001-0003-000", "2024年3月5日", "事件", "6.8", "182", "前墙第3管"],
        ["1-WW-FR-001-0004-000", "2024/4/8", "事件", "7.1", "180", "正常"],
        ["1-WW-FR-001-0003-000", "2024年3月5日", "事件", "6.8", "182", "前墙第3管"],
        ["1-WW-FR-001-0011-000", "2025-02-30", "事件", "7.02", "178", "日期无效"],
    ]
    trs = []
    for row in [header] + data:
        tcs = "".join(cell(v) for v in row)
        trs.append("<w:tr>" + tcs + "</w:tr>")
    body = ("<w:p><w:r><w:t>炉管检测台账（2024）</w:t></w:r></w:p>"
            "<w:tbl>" + "".join(trs) + "</w:tbl>"
            "<w:p><w:r><w:t>以上记录来自现场检验报告。</w:t></w:r></w:p>")
    document = ("<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
                "<w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\">"
                "<w:body>" + body + "<w:sectPr/></w:body></w:document>")
    content_types = ("<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
                     "<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\">"
                     "<Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/>"
                     "<Default Extension=\"xml\" ContentType=\"application/xml\"/>"
                     "<Override PartName=\"/word/document.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml\"/>"
                     "</Types>")
    rels = ("<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
            "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
            "<Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" Target=\"word/document.xml\"/>"
            "</Relationships>")
    path = os.path.join(HERE, "检测台账.docx")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/document.xml", document)
    print("written", path)


if __name__ == "__main__":
    build_events_csv()
    build_material_csv()
    build_unrelated_csv()
    build_xlsx()
    build_docx()
    # openpyxl 读回校验：确保 XLSX 夹具真实可读
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(HERE, "检测记录.xlsx"))
    assert wb.sheetnames == ["测厚记录", "材料库"], wb.sheetnames
    assert wb["测厚记录"].max_row == 5, wb["测厚记录"].max_row
    print("xlsx round-trip ok")
    with zipfile.ZipFile(os.path.join(HERE, "检测台账.docx")) as zf:
        assert "word/document.xml" in zf.namelist()
    print("docx round-trip ok")
